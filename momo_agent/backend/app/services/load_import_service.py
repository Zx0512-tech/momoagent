from __future__ import annotations

import csv
import io
import json
import math
from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.services.load_format_readers import (
    LoadFormatError,
    is_self_describing_format,
    parse_self_describing,
)
from app.services.load_mapping_inference import (
    TRAFFIC_MATRIX_COLUMN_TEMPLATE,
    infer_earthquake_mapping,
    traffic_matrix_columns,
)


MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_ROWS = 200_000
MAX_XLSX_EXTRACTED_BYTES = 100 * 1024 * 1024
MAX_XLSX_MEMBERS = 10_000
# .dat 是地震记录里很常见的纯文本导出后缀（OpenSees 示例、各家台网都用），
# 内容与 .txt 无区别，按同一条通用表格路径读。
TABLE_SUFFIXES = frozenset({'.csv', '.txt', '.dat', '.xlsx'})
# 自描述格式解析后展开成的两列表格列名，供 mapping 直接引用。
SELF_DESCRIBING_TIME_COLUMN = 'time_s'
SELF_DESCRIBING_VALUE_COLUMN = 'acceleration'

# 同一单位的常见写法收敛到换算表的键。上传文件的单位写法五花八门
# （m/s²、m/sec2、CM/S2、Gal），归一化后再查表，避免每加一种写法都要
# 往换算表里塞一条重复条目。
UNIT_ALIASES = {
    'm/s²': 'm/s2',
    'm/s^2': 'm/s2',
    'm/sec2': 'm/s2',
    'm/sec^2': 'm/s2',
    'm/s2': 'm/s2',
    'cm/s²': 'cm/s2',
    'cm/s^2': 'cm/s2',
    'cm/sec2': 'cm/s2',
    'cm/sec^2': 'cm/s2',
    'cm/s2': 'cm/s2',
    'gal': 'cm/s2',
    'mm/s²': 'mm/s2',
    'mm/s^2': 'mm/s2',
    'mm/sec2': 'mm/s2',
    'mm/s2': 'mm/s2',
    'g': 'g',
    'n': 'N',
    'kn': 'kN',
}
# 标准单位固定为 N 与 m/s2。g 用 CGPM 定义的标准重力 9.80665 m/s2，
# 与地震记录（PEER 等）声明 g 时的口径一致。
UNIT_CONVERSIONS = {
    ('FORCE', 'N'): (1.0, 'N'),
    ('FORCE', 'kN'): (1000.0, 'N'),
    ('ACCELERATION', 'm/s2'): (1.0, 'm/s2'),
    ('ACCELERATION', 'g'): (9.80665, 'm/s2'),
    ('ACCELERATION', 'cm/s2'): (0.01, 'm/s2'),
    ('ACCELERATION', 'mm/s2'): (0.001, 'm/s2'),
}


@dataclass(frozen=True)
class StandardizedLoad:
    content: bytes
    digest: str
    report: dict[str, Any]
    # 车流矩阵制品单独存在是不可解释的：还要一份逐节点 mapping 声明哪一列对应哪个
    # 节点。它由标准化时同一次列解析产出，跟着结果一起返回，由调用方登记成制品。
    point_mapping: bytes | None = None


class LoadImportService:
    def inspect(self, file_name: str, content: bytes) -> tuple[dict[str, Any], list[dict[str, str]]]:
        if not content:
            self._error('EMPTY_FILE', '荷载文件为空')
        if len(content) > MAX_UPLOAD_BYTES:
            self._error('FILE_TOO_LARGE', 'MVP 单文件上限为 20 MB')
        suffix = Path(file_name).suffix.lower()
        if suffix not in TABLE_SUFFIXES and not is_self_describing_format(file_name):
            self._error(
                'UNSUPPORTED_FILE_TYPE',
                '仅支持 CSV、TXT、DAT、XLSX 表格和 PEER NGA（.AT1/.AT2）记录文件',
            )

        columns, rows, metadata = self._read_table(suffix, content, file_name=file_name)
        if not rows:
            self._error('NO_DATA_ROWS', '文件中没有数据行')
        if len(rows) > MAX_ROWS:
            self._error('TOO_MANY_ROWS', f'MVP 最多处理 {MAX_ROWS} 行数据')

        profiles = []
        for column in columns:
            values = [row.get(column, '').strip() for row in rows]
            numeric = [number for value in values if (number := self._number_or_none(value)) is not None]
            profiles.append({
                'name': column,
                'numericCount': len(numeric),
                'missingCount': sum(not value for value in values),
                'min': min(numeric) if numeric else None,
                'max': max(numeric) if numeric else None,
                'timeCandidate': column.strip().lower() in {'t', 'time', 'time_s', '时间', '时间(s)'},
            })
        inspection = {
            'fileName': file_name,
            'format': metadata.pop('formatName', None) or suffix.removeprefix('.').upper(),
            'rowCount': len(rows),
            'columnCount': len(columns),
            'columns': profiles,
            'sampleRows': rows[:8],
            **metadata,
        }
        # 建议映射只是给审批界面预填的候选，不参与冻结也不自动执行；真正
        # 生效的 mapping 仍由调用方显式提交，走既有的 SHA256 冻结流程。
        inspection['suggestedMapping'] = infer_earthquake_mapping(
            inspection,
            rows,
            file_name=file_name,
        ).to_dict()
        return inspection, rows

    def standardize(
        self,
        *,
        file_name: str,
        content: bytes,
        mapping: dict[str, Any],
    ) -> StandardizedLoad:
        if mapping.get('version') == 2 or mapping.get('channels'):
            return self._standardize_channels(file_name=file_name, content=content, mapping=mapping)
        inspection, rows = self.inspect(file_name, content)
        columns = {column['name'] for column in inspection['columns']}
        value_column = mapping['valueColumn']
        time_column = mapping.get('timeColumn')
        if value_column not in columns:
            self._error('UNKNOWN_VALUE_COLUMN', f'数值列 {value_column} 不存在')
        if time_column is not None and time_column not in columns:
            self._error('UNKNOWN_TIME_COLUMN', f'时间列 {time_column} 不存在')
        if time_column is None and mapping.get('timeStepS') is None:
            self._error('TIME_STEP_REQUIRED', '无时间列时必须提供 timeStepS')

        times: list[float] = []
        values: list[float] = []
        time_factor = 0.001 if mapping.get('timeUnit') == 'ms' else 1.0
        for index, row in enumerate(rows):
            time_value = index * float(mapping['timeStepS']) if time_column is None else self._required_number(
                row.get(time_column, ''), index + 2, time_column
            ) * time_factor
            value = self._required_number(row.get(value_column, ''), index + 2, value_column)
            times.append(time_value)
            values.append(value)
        if any(current <= previous for previous, current in zip(times, times[1:])):
            self._error('TIME_NOT_STRICTLY_INCREASING', '时间必须严格递增，不能重复或倒序')

        source_unit = mapping['sourceUnit']
        normalized_source_unit = self.normalize_unit(source_unit)
        factor, standard_unit = self._unit_conversion(mapping['quantity'], source_unit)
        output = io.StringIO(newline='')
        writer = csv.writer(output, lineterminator='\n')
        writer.writerow(['time_s', 'target_type', 'target_id', 'component', 'quantity', 'value', 'unit'])
        for time_value, value in zip(times, values):
            writer.writerow([
                self._format_number(time_value),
                mapping['targetType'],
                mapping['targetId'],
                mapping['component'],
                mapping['quantity'],
                self._format_number(value * factor),
                standard_unit,
            ])
        raw = output.getvalue().encode('utf-8')
        digest = sha256(raw).hexdigest()
        report = {
            'sourceFileName': file_name,
            'sourceSha256': sha256(content).hexdigest(),
            'standardSha256': digest,
            'sampleCount': len(rows),
            'timeStartS': times[0],
            'timeEndS': times[-1],
            'sourceTimeUnit': mapping.get('timeUnit', 's') if time_column is not None else 'generated_seconds',
            'sourceEncoding': inspection.get('encoding'),
            'sourceDelimiter': inspection.get('delimiter'),
            'standardEncoding': 'utf-8',
            'standardDelimiter': ',',
            'sourceUnit': source_unit,
            'normalizedSourceUnit': normalized_source_unit,
            'standardUnit': standard_unit,
            'conversionFactor': factor,
            'mapping': mapping,
            'validation': {'nanCount': 0, 'timeStrictlyIncreasing': True},
        }
        return StandardizedLoad(content=raw, digest=digest, report=report)

    def _standardize_channels(
        self,
        *,
        file_name: str,
        content: bytes,
        mapping: dict[str, Any],
    ) -> StandardizedLoad:
        inspection, rows = self.inspect(file_name, content)
        columns = {column['name'] for column in inspection['columns']}
        time_mapping = dict(mapping.get('time') or {})
        time_column = time_mapping.get('column')
        time_step_s = time_mapping.get('stepS')
        if time_column is not None and time_column not in columns:
            self._error('UNKNOWN_TIME_COLUMN', f'时间列 {time_column} 不存在')
        if time_column is None and time_step_s is None:
            self._error('TIME_STEP_REQUIRED', '无时间列时必须提供 time.stepS')
        channels = list(mapping.get('channels') or [])
        # 矩阵型工况（车流）的“数值列”是合成名，不是真实列，所以要在列存在性
        # 检查之前分派。inspection 与 rows 一并传下去，避免对几 MB 的稠密矩阵
        # 再解析一遍。
        if any(channel.get('applicationType') == 'NODAL_FORCE_MATRIX' for channel in channels):
            return self._standardize_matrix(
                file_name=file_name,
                content=content,
                mapping=mapping,
                inspection=inspection,
                rows=rows,
            )
        for channel in channels:
            if channel['valueColumn'] not in columns:
                self._error('UNKNOWN_VALUE_COLUMN', f'数值列 {channel["valueColumn"]} 不存在')

        time_factor = 0.001 if time_mapping.get('unit') == 'ms' else 1.0
        times = [
            index * float(time_step_s)
            if time_column is None
            else self._required_number(row.get(time_column, ''), index + 2, time_column) * time_factor
            for index, row in enumerate(rows)
        ]
        if any(current <= previous for previous, current in zip(times, times[1:])):
            self._error('TIME_NOT_STRICTLY_INCREASING', '时间必须严格递增，不能重复或倒序')

        output = io.StringIO(newline='')
        writer = csv.writer(output, lineterminator='\n')
        writer.writerow([
            'time_s',
            'load_kind',
            'channel_id',
            'application_type',
            'target_type',
            'target_id',
            'component',
            'quantity',
            'value',
            'unit',
        ])
        channel_reports = []
        for channel_index, channel in enumerate(channels, start=1):
            factor, standard_unit = self._unit_conversion(channel['quantity'], channel['sourceUnit'])
            scale = float(channel.get('scale', 1.0))
            channel_id = f'channel_{channel_index}'
            channel_reports.append({
                'channelId': channel_id,
                'sourceUnit': channel['sourceUnit'],
                'normalizedSourceUnit': self.normalize_unit(channel['sourceUnit']),
                'standardUnit': standard_unit,
                'conversionFactor': factor,
                'scale': scale,
                **channel,
            })
            for row_index, (time_value, row) in enumerate(zip(times, rows), start=2):
                value = self._required_number(row.get(channel['valueColumn'], ''), row_index, channel['valueColumn'])
                writer.writerow([
                    self._format_number(time_value),
                    mapping['loadKind'],
                    channel_id,
                    channel['applicationType'],
                    channel.get('targetType') or '',
                    channel.get('targetId') or '',
                    channel['component'],
                    channel['quantity'],
                    self._format_number(value * factor * scale),
                    standard_unit,
                ])
        raw = output.getvalue().encode('utf-8')
        digest = sha256(raw).hexdigest()
        report = {
            'schemaVersion': 2,
            'sourceFileName': file_name,
            'sourceSha256': sha256(content).hexdigest(),
            'standardSha256': digest,
            'loadKind': mapping['loadKind'],
            'sampleCount': len(rows),
            'channelCount': len(channels),
            'outputRowCount': len(rows) * len(channels),
            'timeStartS': times[0],
            'timeEndS': times[-1],
            'sourceEncoding': inspection.get('encoding'),
            'sourceDelimiter': inspection.get('delimiter'),
            'standardEncoding': 'utf-8',
            'standardDelimiter': ',',
            'channels': channel_reports,
            'validation': {'nanCount': 0, 'timeStrictlyIncreasing': True},
        }
        return StandardizedLoad(content=raw, digest=digest, report=report)

    def _standardize_matrix(
        self,
        *,
        file_name: str,
        content: bytes,
        mapping: dict[str, Any],
        inspection: dict[str, Any],
        rows: list[dict[str, str]],
    ) -> StandardizedLoad:
        """把逐节点车流表标准化成稠密矩阵，并同时产出逐节点 mapping。

        与长表分支的区别是产物形态：执行侧 _apply_agent_standard_traffic_load 要
        `time_s` + 每节点一列的矩阵（ANSYS 的 *VREAD 按列取数，OpenSees 的逐节点
        mapping 也按列取），写成长表会丢掉空间分布。

        逐节点绑定不经 HTTP 传递：这里按 traffic_matrix_columns 从源文件表头重新
        解析节点号，与推断时用的是同一条正则、同一份字节，所以两次 standardize()
        （算 expectedSha256 与执行时校验）必然得到同一个绑定，digest 稳定。列名
        统一改写成执行侧要求的 node_{id}_fy_N 口径，源文件写 kN 或不写单位都不影响。
        """

        channels = list(mapping.get('channels') or [])
        if len(channels) != 1:
            self._error('UNSUPPORTED_MATRIX_MAPPING', '矩阵型荷载只支持单通道映射')
        channel = channels[0]
        if channel.get('quantity') != 'FORCE':
            self._error('UNSUPPORTED_MATRIX_MAPPING', '矩阵型荷载通道必须是力时程')

        time_mapping = dict(mapping.get('time') or {})
        time_column = time_mapping.get('column')
        header = [str(profile['name']) for profile in (inspection.get('columns') or [])]
        pairs = traffic_matrix_columns(header, time_column)
        if not pairs:
            self._error(
                'UNKNOWN_MATRIX_COLUMNS',
                '没有 node_<节点号>_fy 形态的逐节点列，或同一节点出现了多列',
            )
        declared_count = channel.get('matrixColumnCount')
        if declared_count is not None and int(declared_count) != len(pairs):
            # 冻结的通道声明了列数，实际解析出的列数不同即源文件已变，必须失败关闭。
            self._error(
                'MATRIX_COLUMN_COUNT_MISMATCH',
                f'映射声明 {int(declared_count)} 个节点列，源文件解析出 {len(pairs)} 个',
            )

        time_step_s = time_mapping.get('stepS')
        time_factor = 0.001 if time_mapping.get('unit') == 'ms' else 1.0
        times = [
            index * float(time_step_s)
            if time_column is None
            else self._required_number(row.get(time_column, ''), index + 2, time_column) * time_factor
            for index, row in enumerate(rows)
        ]
        if any(current <= previous for previous, current in zip(times, times[1:])):
            self._error('TIME_NOT_STRICTLY_INCREASING', '时间必须严格递增，不能重复或倒序')

        factor, standard_unit = self._unit_conversion(channel['quantity'], channel['sourceUnit'])
        scale = float(channel.get('scale', 1.0))
        matrix_columns = [TRAFFIC_MATRIX_COLUMN_TEMPLATE.format(node=node) for node, _ in pairs]
        output = io.StringIO(newline='')
        writer = csv.writer(output, lineterminator='\n')
        writer.writerow(['time_s', *matrix_columns])
        for row_index, (time_value, row) in enumerate(zip(times, rows), start=2):
            writer.writerow([
                self._format_number(time_value),
                *(
                    self._format_number(
                        self._required_number(row.get(source, ''), row_index, source) * factor * scale
                    )
                    for _, source in pairs
                ),
            ])
        raw = output.getvalue().encode('utf-8')
        digest = sha256(raw).hexdigest()

        # source_column 是矩阵里的 1 基列号（时间列不计），与上面写出的列顺序同源。
        node_mappings = [
            {
                'fem_node_id': node,
                'dof': 'FY',
                'scale': 1.0,
                'source_column': index,
                'group': 'traffic_deck',
                'label': f'traffic_deck_{node}',
            }
            for index, (node, _) in enumerate(pairs, start=1)
        ]
        point_mapping = json.dumps(
            node_mappings,
            ensure_ascii=False,
            separators=(',', ':'),
            allow_nan=False,
        ).encode('utf-8')

        target_nodes = [node for node, _ in pairs]
        report = {
            'schemaVersion': 2,
            'sourceFileName': file_name,
            'sourceSha256': sha256(content).hexdigest(),
            'standardSha256': digest,
            'loadKind': mapping['loadKind'],
            'applicationType': 'NODAL_FORCE_MATRIX',
            'sampleCount': len(rows),
            'channelCount': 1,
            'nodeCount': len(pairs),
            'outputRowCount': len(rows),
            'timeStartS': times[0],
            'timeEndS': times[-1],
            'sourceEncoding': inspection.get('encoding'),
            'sourceDelimiter': inspection.get('delimiter'),
            'standardEncoding': 'utf-8',
            'standardDelimiter': ',',
            'targetSetId': channel.get('targetId'),
            'targetNodes': target_nodes,
            'component': channel.get('component'),
            'sourceUnit': channel['sourceUnit'],
            'normalizedSourceUnit': self.normalize_unit(channel['sourceUnit']),
            'standardUnit': standard_unit,
            'conversionFactor': factor,
            'scale': scale,
            'sourceColumns': [source for _, source in pairs],
            'matrixColumns': matrix_columns,
            'pointMappingSha256': sha256(point_mapping).hexdigest(),
            # 逐节点独立时程，求解侧不得再做等权分配或求和。
            'distribution': 'PER_NODE_INDEPENDENT_TIME_HISTORY',
            'channels': [{
                'channelId': 'channel_1',
                'standardUnit': standard_unit,
                'conversionFactor': factor,
                **channel,
            }],
            'mapping': mapping,
            'validation': {'nanCount': 0, 'timeStrictlyIncreasing': True},
        }
        return StandardizedLoad(
            content=raw,
            digest=digest,
            report=report,
            point_mapping=point_mapping,
        )

    def _read_self_describing(
        self,
        file_name: str,
        content: bytes,
    ) -> tuple[list[str], list[dict[str, str]], dict[str, Any]]:
        """把自描述格式展开成两列表格，并带出文件头声明的元数据。

        展开成 time_s + acceleration 两列后，下游 standardize 与通用表格
        完全同路——时间列真实存在，dt 与单位来自文件头而非推断。
        """

        try:
            record = parse_self_describing(file_name, content)
        except LoadFormatError as exc:
            self._error(exc.code, exc.message)
        columns = [SELF_DESCRIBING_TIME_COLUMN, SELF_DESCRIBING_VALUE_COLUMN]
        rows = [
            {
                SELF_DESCRIBING_TIME_COLUMN: self._format_number(index * record.dt),
                SELF_DESCRIBING_VALUE_COLUMN: self._format_number(value),
            }
            for index, value in enumerate(record.values)
        ]
        metadata = {
            'sheetName': None,
            'encoding': 'utf-8',
            'delimiter': None,
            'formatName': record.format_name,
            # selfDescribing 让推断层知道 dt/单位是读出来的，可以直接采信。
            'selfDescribing': {
                'formatName': record.format_name,
                'timeColumn': SELF_DESCRIBING_TIME_COLUMN,
                'valueColumn': SELF_DESCRIBING_VALUE_COLUMN,
                'dt': record.dt,
                'sourceUnit': record.unit,
                'quantity': record.quantity,
                'sampleCount': record.sample_count,
                'durationS': record.duration_s,
                'headerLines': list(record.header_lines),
                **{key: value for key, value in record.metadata.items() if key != 'declaredDt'},
            },
        }
        return columns, rows, metadata

    def _read_table(
        self,
        suffix: str,
        content: bytes,
        *,
        file_name: str,
    ) -> tuple[list[str], list[dict[str, str]], dict[str, Any]]:
        if is_self_describing_format(file_name):
            return self._read_self_describing(file_name, content)
        if suffix == '.xlsx':
            self._validate_xlsx_archive(content)
            workbook = None
            try:
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                sheet = workbook[workbook.sheetnames[0]]
                sheet_name = sheet.title
                raw_rows = [[self._cell_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError, EOFError):
                self._error('INVALID_XLSX', 'XLSX 文件损坏或结构不合法')
            finally:
                if workbook is not None:
                    workbook.close()
            columns, rows = self._rows_to_records(raw_rows, has_header=True)
            return columns, rows, {'sheetName': sheet_name, 'encoding': None, 'delimiter': None}

        text, encoding = self._decode_text(content)
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=',;\t ')
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ','
        parsed = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        parsed = [row for row in parsed if any(cell.strip() for cell in row)]
        # 地震记录常见"数值用连续空格/制表符对齐"的写法，csv.Sniffer 对这类
        # 文件会失败并退回逗号，整行被当成一个非数值单元格。按空白串重切能
        # 得到更多列时改用它——多列一定比整行一格更接近真实结构。
        # 只比数据块的宽度：文件开头的说明文字词数可能比真实列数还多，按全表
        # 最宽行比较会把带前言的逗号文件误判成空白分隔，整块数据挤成一列。
        whitespace_parsed = _split_on_whitespace_runs(text)
        if _data_block_width(whitespace_parsed) > _data_block_width(parsed):
            parsed = whitespace_parsed
            delimiter = 'whitespace'
        parsed, preamble_lines = _split_preamble(parsed)
        if not parsed:
            self._error('EMPTY_FILE', '荷载文件为空')
        # 首行是不是表头，只看"整行是否全为数值"。csv.Sniffer().has_header 对
        # 单列纯数值文件会误判成有表头，把第一个采样点当列名吃掉——输出的样本
        # 数、等步长、峰值全都合法，只是整条时程少一点且平移了一个 dt。
        has_header = _all_numbers_or_none(' '.join(parsed[0])) is None
        columns, rows = self._rows_to_records(parsed, has_header=has_header)
        return columns, rows, {
            'sheetName': None,
            'encoding': encoding,
            'delimiter': delimiter,
            'preambleLines': preamble_lines,
        }

    @staticmethod
    def _validate_xlsx_archive(content: bytes) -> None:
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                members = archive.infolist()
                if len(members) > MAX_XLSX_MEMBERS:
                    LoadImportService._error('XLSX_ARCHIVE_TOO_LARGE', 'XLSX 内部文件数量超出限制')
                if sum(member.file_size for member in members) > MAX_XLSX_EXTRACTED_BYTES:
                    LoadImportService._error('XLSX_ARCHIVE_TOO_LARGE', 'XLSX 解压后大小超出限制')
                if any(member.flag_bits & 0x1 for member in members):
                    LoadImportService._error('ENCRYPTED_XLSX', '不支持加密 XLSX 文件')
        except BadZipFile:
            LoadImportService._error('INVALID_XLSX', 'XLSX 文件损坏或结构不合法')

    def _rows_to_records(
        self,
        raw_rows: list[list[str]],
        *,
        has_header: bool,
    ) -> tuple[list[str], list[dict[str, str]]]:
        if not raw_rows:
            self._error('EMPTY_FILE', '荷载文件为空')
        width = max(len(row) for row in raw_rows)
        if has_header:
            columns = [cell.strip() or f'column_{index + 1}' for index, cell in enumerate(raw_rows[0])]
            data_rows = raw_rows[1:]
        else:
            columns = [f'column_{index + 1}' for index in range(width)]
            data_rows = raw_rows
        if len(columns) != len(set(columns)):
            self._error('DUPLICATE_COLUMNS', '文件包含重复列名')
        records = []
        for row in data_rows:
            padded = [*row, *([''] * max(0, len(columns) - len(row)))]
            records.append({column: padded[index].strip() for index, column in enumerate(columns)})
        return columns, records

    @staticmethod
    def _decode_text(content: bytes) -> tuple[str, str]:
        for encoding in ('utf-8-sig', 'utf-8', 'gb18030'):
            try:
                return content.decode(encoding), encoding
            except UnicodeDecodeError:
                continue
        LoadImportService._error('UNSUPPORTED_ENCODING', '无法识别文本编码')

    @staticmethod
    def normalize_unit(source_unit: str) -> str:
        """把同一单位的常见写法收敛成换算表的键。"""

        text = str(source_unit or '').strip()
        return UNIT_ALIASES.get(text.lower(), text)

    @staticmethod
    def _unit_conversion(quantity: str, source_unit: str) -> tuple[float, str]:
        normalized = LoadImportService.normalize_unit(source_unit)
        try:
            return UNIT_CONVERSIONS[(quantity, normalized)]
        except KeyError:
            LoadImportService._error('INCOMPATIBLE_UNIT', f'{source_unit} 与 {quantity} 不兼容')

    @staticmethod
    def _required_number(value: str, row_number: int, column: str) -> float:
        number = LoadImportService._number_or_none(value)
        if number is None:
            LoadImportService._error('NON_NUMERIC_VALUE', f'第 {row_number} 行 {column} 不是有限数值')
        return number

    @staticmethod
    def _number_or_none(value: str) -> float | None:
        try:
            number = float(value)
        except (TypeError, ValueError):
            return None
        return number if math.isfinite(number) else None

    @staticmethod
    def _format_number(value: float) -> str:
        return format(value, '.15g')

    @staticmethod
    def _cell_text(value: object) -> str:
        return '' if value is None else str(value)

    @staticmethod
    def _error(code: str, message: str) -> None:
        raise HTTPException(status_code=422, detail={'code': code, 'message': message})


def _split_on_whitespace_runs(text: str) -> list[list[str]]:
    """按连续空白切分每一行，丢掉空行。"""

    rows: list[list[str]] = []
    for line in text.splitlines():
        cells = [cell for cell in line.strip().split() if cell]
        if cells:
            rows.append(cells)
    return rows


def _data_block_width(rows: list[list[str]]) -> int:
    """数据块里最宽一行的列数——只算整行全为数值的行。没有数据行时为 0。"""

    return max(
        (len(row) for row in rows if _all_numbers_or_none(' '.join(row)) is not None),
        default=0,
    )


def _split_preamble(rows: list[list[str]]) -> tuple[list[list[str]], list[str]]:
    """把数据块之前的说明文字摘掉，只留表头行和数据行。

    真实记录常在数据前写几行事件说明（来源、台站、频段），紧挨数据块的那一
    行才可能是表头。判据用列数：表头的列数与数据块一致，说明行通常不一致；
    列数对不上就连它一起摘掉，用合成列名，免得拿"Frequency range:"这种词
    当列名。摘掉的行会随 inspection 返回，不静默丢弃用户文件里的内容。
    """

    first_data = next(
        (
            index
            for index, row in enumerate(rows)
            if _all_numbers_or_none(' '.join(row)) is not None
        ),
        None,
    )
    if not first_data:
        # 没有数据行，或首行就是数据（无前言无表头），两种都原样交给下游。
        return rows, []
    width = _data_block_width(rows[first_data:])
    header = rows[first_data - 1]
    if len(header) != width:
        # 表头行的分隔符可能与数据块不同（数据用制表符对齐、表头用空格），
        # 按空白重切一次。宽度对上就仍当表头用：写着 "Accel[g]" 的表头一旦
        # 被当说明文字丢掉，单位就只能靠峰值量级猜，而文件里本来是声明了的。
        respaced = ' '.join(header).split()
        header = respaced if len(respaced) == width else header
    if len(header) == width:
        return [header, *rows[first_data:]], [
            ' '.join(row).strip() for row in rows[:first_data - 1]
        ]
    return rows[first_data:], [' '.join(row).strip() for row in rows[:first_data]]


def _all_numbers_or_none(line: str) -> list[float] | None:
    """整行都是有限数值时返回解析结果，否则返回 None。"""

    cells = [cell for cell in line.replace(',', ' ').split() if cell]
    if not cells:
        return None
    numbers: list[float] = []
    for cell in cells:
        number = LoadImportService._number_or_none(cell)
        if number is None:
            return None
        numbers.append(number)
    return numbers


load_import_service = LoadImportService()
