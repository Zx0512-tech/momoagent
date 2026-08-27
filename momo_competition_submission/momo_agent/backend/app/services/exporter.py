from csv import writer
from io import BytesIO, StringIO
from pathlib import Path
import sys
import tempfile

REPO_ROOT = Path(__file__).resolve().parents[4]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from pyansys_bridge.models import TimeHistoryLoad


def build_export_manifest(project_name: str) -> list[dict[str, str]]:
    base_name = project_name or "bridge-wind"
    return [
        {"name": f"{base_name}-wind-results.xlsx", "type": "excel"},
        {"name": f"{base_name}-girder-velocity.csv", "type": "csv"},
        {"name": f"{base_name}-tower-load.csv", "type": "csv"},
        {"name": f"{base_name}-girder-all-points.csv", "type": "csv"},
        {"name": f"{base_name}-tower-all-points.csv", "type": "csv"},
        {"name": f"{base_name}-spectrum.csv", "type": "csv"},
        {"name": f"{base_name}-coherence.csv", "type": "csv"},
    ]



def build_excel_bytes(result: dict) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "计算摘要"
    summary_sheet.append(["项目", result["summary"]["projectName"]])
    summary_sheet.append(["主梁 Ud", result["summary"]["girder"]["ud"]])
    summary_sheet.append(["主梁 Iu", result["summary"]["girder"]["iu"]])
    summary_sheet.append(["主梁 Iv", result["summary"]["girder"]["iv"]])
    summary_sheet.append(["主梁 Iw", result["summary"]["girder"]["iw"]])
    summary_sheet.append(["主梁控制点数", result["summary"]["girder"]["pointCount"]])
    summary_sheet.append(["均值误差", result["summary"]["girder"]["validation"]["meanError"]])
    summary_sheet.append(["标准差误差", result["summary"]["girder"]["validation"]["stdError"]])
    summary_sheet.append(["主梁全点最大均值误差", result["summary"]["girder"]["pointValidation"]["maxAbsMeanError"]])
    summary_sheet.append(["主梁全点最大标准差误差", result["summary"]["girder"]["pointValidation"]["maxAbsStdError"]])
    summary_sheet.append(["桥塔分层数", result["summary"]["tower"]["levelCount"]])
    summary_sheet.append(["桥塔顶部 Ud", result["summary"]["tower"]["topUd"]])
    summary_sheet.append(["桥塔全点最大均值误差", result["summary"]["tower"]["pointValidation"]["maxAbsMeanError"]])
    summary_sheet.append(["桥塔全点最大标准差误差", result["summary"]["tower"]["pointValidation"]["maxAbsStdError"]])
    summary_sheet.append(["主梁相干平均误差", result["coherenceValidation"]["girder"]["meanAbsError"]])
    summary_sheet.append(["主梁相干最大误差", result["coherenceValidation"]["girder"]["maxAbsError"]])
    summary_sheet.append(["桥塔相干平均误差", result["coherenceValidation"]["tower"]["meanAbsError"]])
    summary_sheet.append(["桥塔相干最大误差", result["coherenceValidation"]["tower"]["maxAbsError"]])
    summary_sheet.append(["主梁风速峰值因子最大值", result["extremeValidation"]["peakFactor"]["girderVelocity"]["max"]])
    summary_sheet.append(["桥塔风速峰值因子最大值", result["extremeValidation"]["peakFactor"]["towerVelocity"]["max"]])

    validation_sheet = workbook.create_sheet("校核统计")
    validation_sheet.append(["类别", "指标", "min", "max", "mean", "std", "p05", "p95"])
    for name, values in result["extremeValidation"]["extremes"].items():
        validation_sheet.append([
            name,
            "extreme",
            values["min"],
            values["max"],
            values["mean"],
            values["std"],
            values["p05"],
            values["p95"],
        ])
    validation_sheet.append(["主梁风速", "peakFactor", result["extremeValidation"]["peakFactor"]["girderVelocity"]["min"], result["extremeValidation"]["peakFactor"]["girderVelocity"]["max"], result["extremeValidation"]["peakFactor"]["girderVelocity"]["mean"], "", "", ""])
    validation_sheet.append(["桥塔风速", "peakFactor", result["extremeValidation"]["peakFactor"]["towerVelocity"]["min"], result["extremeValidation"]["peakFactor"]["towerVelocity"]["max"], result["extremeValidation"]["peakFactor"]["towerVelocity"]["mean"], "", "", ""])

    spectrum_sheet = workbook.create_sheet("风谱")
    spectrum_sheet.append(["序列", "频率", "谱值"])
    for series in result["spectra"]:
        for frequency, value in zip(series["frequency"], series["values"], strict=False):
            spectrum_sheet.append([series["name"], frequency, value])

    coherence_sheet = workbook.create_sheet("相干函数")
    coherence_sheet.append(["序列", "频率", "相干值"])
    for series in result["coherence"]:
        for frequency, value in zip(series["frequency"], series["values"], strict=False):
            coherence_sheet.append([series["name"], frequency, value])

    girder_sheet = workbook.create_sheet("主梁时程")
    girder_sheet.append(["时间", "风速", "水平力", "竖向力", "力矩"])
    histories = result["histories"]
    for time_value, velocity, fh, fv, moment in zip(
        histories["time"],
        histories["girder"]["velocity"],
        histories["girder"]["fh"],
        histories["girder"]["fv"],
        histories["girder"]["m"],
        strict=False,
    ):
        girder_sheet.append([time_value, velocity, fh, fv, moment])

    tower_sheet = workbook.create_sheet("桥塔时程")
    tower_sheet.append(["时间", "风速", "拖曳力"])
    for time_value, velocity, drag in zip(
        histories["time"],
        histories["tower"]["velocity"],
        histories["tower"]["drag"],
        strict=False,
    ):
        tower_sheet.append([time_value, velocity, drag])

    girder_all_sheet = workbook.create_sheet("主梁全点时程")
    girder_all_sheet.append(["控制点", "时间", "风速", "水平力", "竖向力", "力矩"])
    write_girder_point_rows(girder_all_sheet.append, result)

    tower_all_sheet = workbook.create_sheet("桥塔全点时程")
    tower_all_sheet.append(["高度", "时间", "风速", "拖曳力"])
    write_tower_point_rows(tower_all_sheet.append, result)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()



def build_girder_csv_bytes(result: dict) -> bytes:
    stream = StringIO(newline="")
    csv_writer = writer(stream)
    csv_writer.writerow(["时间", "风速", "水平力", "竖向力", "力矩"])
    histories = result["histories"]
    for time_value, velocity, fh, fv, moment in zip(
        histories["time"],
        histories["girder"]["velocity"],
        histories["girder"]["fh"],
        histories["girder"]["fv"],
        histories["girder"]["m"],
        strict=False,
    ):
        csv_writer.writerow([time_value, velocity, fh, fv, moment])
    return stream.getvalue().encode("utf-8-sig")


def build_unified_wind_csv_bytes(result: dict) -> bytes:
    histories = result["histories"]
    time_values = histories["time"]
    if len(time_values) < 2:
        raise ValueError("wind time history must contain at least two samples")
    dt = float(time_values[1] - time_values[0])
    load = TimeHistoryLoad.from_series(
        kind="wind",
        samples=tuple(float(value) for value in histories["girder"]["fh"]),
        dt=dt,
        duration=float(time_values[-1]),
        application={"type": "nodal_force", "component": "girder_fh", "dof": "FX"},
        metadata={
            "name": result["summary"]["projectName"] or "bridge-wind",
            "unit": "N",
            "source": "momo_agent.histories.girder.fh",
        },
    )
    with tempfile.TemporaryDirectory() as temp_dir:
        path = Path(temp_dir) / "unified_wind.csv"
        load.write_csv(path)
        return path.read_bytes()


def build_tower_csv_bytes(result: dict) -> bytes:
    stream = StringIO(newline="")
    csv_writer = writer(stream)
    csv_writer.writerow(["时间", "风速", "拖曳力"])
    histories = result["histories"]
    for time_value, velocity, drag in zip(
        histories["time"],
        histories["tower"]["velocity"],
        histories["tower"]["drag"],
        strict=False,
    ):
        csv_writer.writerow([time_value, velocity, drag])
    return stream.getvalue().encode("utf-8-sig")


def write_girder_point_rows(write_row, result: dict) -> None:
    histories = result["histories"]
    points = result["pointHistories"]["girder"]["points"]
    point_histories = result["pointHistories"]["girder"]
    for point, velocity_row, fh_row, fv_row, moment_row in zip(
        points,
        point_histories["velocity"],
        point_histories["fh"],
        point_histories["fv"],
        point_histories["m"],
        strict=False,
    ):
        for time_value, velocity, fh, fv, moment in zip(
            histories["time"],
            velocity_row,
            fh_row,
            fv_row,
            moment_row,
            strict=False,
        ):
            write_row([point, time_value, velocity, fh, fv, moment])


def write_tower_point_rows(write_row, result: dict) -> None:
    histories = result["histories"]
    points = result["pointHistories"]["tower"]["points"]
    point_histories = result["pointHistories"]["tower"]
    for point, velocity_row, drag_row in zip(
        points,
        point_histories["velocity"],
        point_histories["drag"],
        strict=False,
    ):
        for time_value, velocity, drag in zip(histories["time"], velocity_row, drag_row, strict=False):
            write_row([point, time_value, velocity, drag])


def build_girder_all_points_csv_bytes(result: dict) -> bytes:
    stream = StringIO(newline="")
    csv_writer = writer(stream)
    csv_writer.writerow(["控制点", "时间", "风速", "水平力", "竖向力", "力矩"])
    write_girder_point_rows(csv_writer.writerow, result)
    return stream.getvalue().encode("utf-8-sig")


def build_tower_all_points_csv_bytes(result: dict) -> bytes:
    stream = StringIO(newline="")
    csv_writer = writer(stream)
    csv_writer.writerow(["高度", "时间", "风速", "拖曳力"])
    write_tower_point_rows(csv_writer.writerow, result)
    return stream.getvalue().encode("utf-8-sig")


def build_series_csv_bytes(series_list: list[dict], value_header: str) -> bytes:
    stream = StringIO(newline="")
    csv_writer = writer(stream)
    csv_writer.writerow(["序列", "频率", value_header])
    for series in series_list:
        for frequency, value in zip(series["frequency"], series["values"], strict=False):
            csv_writer.writerow([series["name"], frequency, value])
    return stream.getvalue().encode("utf-8-sig")


def build_spectrum_csv_bytes(result: dict) -> bytes:
    return build_series_csv_bytes(result["spectra"], "谱值")


def build_coherence_csv_bytes(result: dict) -> bytes:
    return build_series_csv_bytes(result["coherence"], "相干值")
